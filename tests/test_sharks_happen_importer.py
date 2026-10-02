from __future__ import annotations

import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.mongodb import COLLECTIONS
from app.services.sharks_happen_importer import (
    SIDE_EFFECTS,
    import_sharks_happen,
    read_workbook,
    source_link_from_sharks_happen_record,
)


NOW = datetime(2026, 10, 2, 12, 0, tzinfo=timezone.utc)
HEADERS = [
    "Count", "Date", "Victim", "Age", "Gender", "Location", "City/State", "Country",
    "Depth of Water in Feet", "Water Temp", "Visibility", "Weather", "Distance from Shore",
    "Distance from Boat in Feet", "Time", "AM/PM", "Shark", "Shark Size", "Injury", "Fatal",
    "Consume / Attempt Consume", "Bitten in Half", "Swallowed Whole", "Depth of Attack (Divers)",
    "Activity", "Clothing", "Weapon", "Eye Gouge", "Seen Submerge", "Bypass Others", "Circle",
    "Urinating", "Pressure", "Humidity", "Cut",
]


class FakeCollection:
    def __init__(self):
        self.replacements = []
        self.inserted = []

    def replace_one(self, query, document, upsert=False):
        self.replacements.append((query, document, upsert))

    def insert_one(self, document):
        self.inserted.append(document)


class FakeDatabase:
    def __init__(self):
        self.collections = {}

    def __getitem__(self, name):
        return self.collections.setdefault(name, FakeCollection())


def create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Sharks Happen Stats"])
    sheet.append([])
    sheet.append(HEADERS)
    sheet.append([1, datetime(2001, 6, 2), "Private Person", 30, "Male", "Example Beach", "Queensland", "Australia", None, None, None, None, None, None, None, None, "GW", "10 Ft", "Leg injury", "Yes", None, None, None, None, "Surfing"])
    sheet.append([2, "12-?-1949", "Second Person", None, "Female", "Other Beach", None, "USA", None, None, None, None, None, None, None, None, None, None, "Minor", None, None, None, None, None, "Swimming"])
    sheet.append([3])
    sheet.append(["Total", None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None])
    workbook.save(path)


def create_comparison_database(path: Path) -> None:
    with sqlite3.connect(path) as connection:
        connection.execute(
            """CREATE TABLE incidents_scrubbed (
                record_id TEXT, canonical_key TEXT, source_name TEXT, source_record_id TEXT,
                date_text TEXT, year INTEGER, month INTEGER, day INTEGER, country TEXT, area TEXT,
                location_public TEXT, activity TEXT, species_common TEXT, fatal INTEGER
            )"""
        )
        connection.execute(
            "INSERT INTO incidents_scrubbed VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            ("candidate-1", "case:example", "gsaf_latest_xls", "2001.06.02", "2-Jun-2001", 2001, 6, 2, "AUSTRALIA", "Queensland", "Example Beach", "surfing", "white shark", 1),
        )


def test_read_workbook_preserves_source_claims_and_unknowns(tmp_path: Path):
    workbook_path = tmp_path / "Sharks Happen Stats.xlsx"
    create_workbook(workbook_path)

    records = read_workbook(workbook_path, imported_at=NOW)

    assert len(records) == 2
    assert records[0]["source_record_id"] == "sharks_happen:0001"
    assert records[0]["victim_name_private"] == "Private Person"
    assert records[0]["fatality_claim"] is True
    assert records[1]["fatality_claim"] is None
    assert records[1]["incident_date_normalized"] is None
    assert records[1]["year"] == 1949
    assert "unparsed_date" in records[1]["normalization_warnings"]
    assert records[0]["side_effects"] == SIDE_EFFECTS


def test_import_reports_duplicate_candidates_without_merging(tmp_path: Path):
    workbook_path = tmp_path / "Sharks Happen Stats.xlsx"
    database_path = tmp_path / "comparison.sqlite"
    staging_path = tmp_path / "staging.json"
    report_path = tmp_path / "report.json"
    review_path = tmp_path / "review.csv"
    create_workbook(workbook_path)
    create_comparison_database(database_path)

    report = import_sharks_happen(
        workbook_path,
        comparison_database=database_path,
        staging_path=staging_path,
        report_path=report_path,
        review_path=review_path,
        imported_at=NOW,
    )

    assert report["records_written"] == 2
    assert report["duplicate_review_counts"]["exact_candidate"] == 1
    assert report["automatic_merges"] == 0
    assert report["automatic_registry_promotions"] == 0
    assert staging_path.exists()
    assert report_path.exists()
    assert review_path.exists()


def test_registry_link_is_attributed_and_private(tmp_path: Path):
    workbook_path = tmp_path / "Sharks Happen Stats.xlsx"
    create_workbook(workbook_path)
    record = read_workbook(workbook_path, imported_at=NOW)[0]

    link = source_link_from_sharks_happen_record(record)

    assert link.source_type == "sharks_happen_spreadsheet"
    assert link.source_name == "Sharks Happen Stats"
    assert link.public_citation_allowed is False
    assert "source_species_label" in link.linked_claims


def test_within_workbook_repeats_are_flagged_without_merging(tmp_path: Path):
    workbook_path = tmp_path / "Sharks Happen Stats.xlsx"
    create_workbook(workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook.active
    duplicate = [3, datetime(2001, 6, 2), "Private Person", 30, "Male", "Example Beach", "Queensland", "Australia", None, None, None, None, None, None, None, None, "GW", "10 Ft", "Leg injury", "Yes", None, None, None, None, "Surfing"]
    sheet.insert_rows(6)
    for column, value in enumerate(duplicate, start=1):
        sheet.cell(6, column, value)
    workbook.save(workbook_path)

    report = import_sharks_happen(
        workbook_path,
        comparison_database=tmp_path / "missing.sqlite",
        staging_path=tmp_path / "staging.json",
        report_path=tmp_path / "report.json",
        review_path=tmp_path / "review.csv",
        imported_at=NOW,
    )

    assert report["records_written"] == 3
    assert report["within_source_possible_duplicate_rows"] == 2
    assert report["automatic_merges"] == 0


def test_mongo_persistence_is_internal_and_separate(tmp_path: Path):
    workbook_path = tmp_path / "Sharks Happen Stats.xlsx"
    create_workbook(workbook_path)
    database = FakeDatabase()

    report = import_sharks_happen(
        workbook_path,
        comparison_database=tmp_path / "missing.sqlite",
        staging_path=tmp_path / "staging.json",
        report_path=tmp_path / "report.json",
        imported_at=NOW,
        mongo_database=database,
    )

    source_collection = database[COLLECTIONS["sharks_happen_sources"]]
    report_collection = database[COLLECTIONS["sharks_happen_import_reports"]]
    assert len(source_collection.replacements) == 2
    assert source_collection.replacements[0][1]["visibility"] == "internal"
    assert len(report_collection.inserted) == 1
    assert report["mongo_persistence"]["records_upserted"] == 2
