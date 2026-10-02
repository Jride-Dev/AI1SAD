from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import date, datetime, time, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.incident_registry import SourceLink


SOURCE_NAME = "Sharks Happen Stats"
SOURCE_CREATOR = "Hal"
SOURCE_CHANNEL = "@sharkshappen on YouTube"
SCHEMA_VERSION = 1

DEFAULT_STAGING_PATH = Path("data/imports/sharks_happen/staging/latest_sharks_happen_sources.json")
DEFAULT_REPORT_PATH = Path("data/imports/sharks_happen/reports/latest_sharks_happen_import_report.json")
DEFAULT_REVIEW_PATH = Path("data/imports/sharks_happen/reports/latest_duplicate_review.csv")
DEFAULT_COMPARISON_DATABASE = Path("data/processed/complete_incidents_scrubbed.sqlite")

SIDE_EFFECTS = {
    "creates_warnings": False,
    "creates_alerts": False,
    "creates_public_feed_entries": False,
    "creates_replay_facts": False,
    "creates_drone_observations": False,
    "alters_scoring": False,
    "alters_replay": False,
    "promotes_registry_cases": False,
    "merges_duplicate_candidates": False,
}

COUNTRY_ALIASES = {
    "SA": "SOUTH AFRICA",
    "SOUTH AFRICA": "SOUTH AFRICA",
    "USA": "USA",
    "US": "USA",
    "UNITED STATES": "USA",
    "AUSTRALIA": "AUSTRALIA",
}

DATE_FORMATS = (
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%m-%d-%Y",
    "%d-%m-%Y",
    "%d-%b-%Y",
    "%d-%B-%Y",
)


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        text = re.sub(r"\s+", " ", value).strip()
        return text or None
    return value


def text_value(value: Any) -> str | None:
    cleaned = clean_cell(value)
    if cleaned is None:
        return None
    text = str(cleaned).strip()
    return text or None


def normalized_text(value: Any) -> str:
    text = (text_value(value) or "").upper()
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalized_country(value: Any) -> str | None:
    key = normalized_text(value)
    return COUNTRY_ALIASES.get(key, key or None)


def parse_date(value: Any) -> tuple[str | None, int | None, int | None, int | None, list[str]]:
    if isinstance(value, datetime):
        return value.date().isoformat(), value.year, value.month, value.day, []
    if isinstance(value, date):
        return value.isoformat(), value.year, value.month, value.day, []
    raw = text_value(value)
    if not raw:
        return None, None, None, None, ["missing_date"]
    for fmt in DATE_FORMATS:
        try:
            parsed = datetime.strptime(raw, fmt).date()
            return parsed.isoformat(), parsed.year, parsed.month, parsed.day, []
        except ValueError:
            continue
    year_match = re.search(r"\b(18|19|20)\d{2}\b", raw)
    return None, int(year_match.group(0)) if year_match else None, None, None, ["unparsed_date"]


def parse_boolean_claim(value: Any) -> bool | None:
    text = normalized_text(value)
    if text in {"YES", "Y", "TRUE", "1"}:
        return True
    if text in {"NO", "N", "FALSE", "0"}:
        return False
    return None


def stable_fingerprint(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, ensure_ascii=True, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def read_workbook(path: str | Path, *, imported_at: datetime | None = None) -> list[dict[str, Any]]:
    input_path = Path(path)
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [text_value(sheet.cell(3, column).value) or f"column_{column}" for column in range(1, sheet.max_column + 1)]
    timestamp = imported_at or utc_now()
    records: list[dict[str, Any]] = []

    for worksheet_row, row in enumerate(
        sheet.iter_rows(min_row=4, max_col=sheet.max_column, values_only=True),
        start=4,
    ):
        values = list(row)
        source_count = values[0]
        if not isinstance(source_count, (int, float)) or isinstance(source_count, bool):
            continue
        source_count = int(source_count)
        if not any(values[index] not in (None, "") for index in (1, 2, 5, 8, 16, 18, 24)):
            continue

        raw_claims = {headers[index]: clean_cell(value) for index, value in enumerate(values)}
        date_normalized, year, month, day, warnings = parse_date(values[1])
        fatality = parse_boolean_claim(values[19])
        consume_attempt = parse_boolean_claim(values[20])
        source_record_id = f"sharks_happen:{source_count:04d}"
        matching_payload = {
            "date": date_normalized,
            "year": year,
            "victim": normalized_text(values[2]),
            "location": normalized_text(values[5]),
            "area": normalized_text(values[6]),
            "country": normalized_country(values[7]),
            "activity": normalized_text(values[24]),
        }
        if fatality is None and text_value(values[19]):
            warnings.append("uncertain_fatality_value")
        if consume_attempt is None and text_value(values[20]):
            warnings.append("uncertain_consumption_value")

        record = {
            "source_record_id": source_record_id,
            "source_name": SOURCE_NAME,
            "source_creator": SOURCE_CREATOR,
            "source_channel": SOURCE_CHANNEL,
            "source_workbook": input_path.name,
            "source_sheet": sheet.title,
            "source_row_number": worksheet_row,
            "source_count": source_count,
            "imported_at": timestamp.isoformat(),
            "visibility": "internal",
            "review_status": "unreviewed_source_claim",
            "victim_name_private": text_value(values[2]),
            "age_raw": text_value(values[3]),
            "gender_raw": text_value(values[4]),
            "incident_date_raw": text_value(values[1]),
            "incident_date_normalized": date_normalized,
            "year": year,
            "month": month,
            "day": day,
            "location_raw": text_value(values[5]),
            "area_raw": text_value(values[6]),
            "country_raw": text_value(values[7]),
            "country_normalized": normalized_country(values[7]),
            "activity_raw": text_value(values[24]),
            "shark_label_raw": text_value(values[16]),
            "shark_size_raw": text_value(values[17]),
            "injury_raw": text_value(values[18]),
            "fatality_claim": fatality,
            "consumption_attempt_claim": consume_attempt,
            "raw_claims": raw_claims,
            "normalization_warnings": sorted(set(warnings)),
            "source_rights_note": "Privately supplied source workbook; redistribution and public release are not authorized by this import.",
            "provenance_note": "Values are source-attributed claims from Hal's Sharks Happen Stats workbook, not AI1SAD-confirmed facts.",
            "side_effects": dict(SIDE_EFFECTS),
        }
        record["source_fingerprint"] = stable_fingerprint(raw_claims)
        record["within_source_match_key"] = stable_fingerprint(matching_payload)[:24]
        records.append(record)
    workbook.close()
    return records


def similarity(left: Any, right: Any) -> float:
    a = normalized_text(left)
    b = normalized_text(right)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    a_tokens = set(a.split())
    b_tokens = set(b.split())
    token_score = len(a_tokens & b_tokens) / len(a_tokens | b_tokens) if a_tokens | b_tokens else 0.0
    return max(token_score, SequenceMatcher(None, a, b).ratio())


def load_comparison_records(path: str | Path) -> list[dict[str, Any]]:
    database_path = Path(path)
    if not database_path.exists():
        return []
    with sqlite3.connect(database_path) as connection:
        connection.row_factory = sqlite3.Row
        return [dict(row) for row in connection.execute("SELECT * FROM incidents_scrubbed")]


def score_candidate(source: dict[str, Any], candidate: dict[str, Any]) -> dict[str, Any]:
    score = 0.0
    matched: list[str] = []
    conflicts: list[str] = []
    source_date = (source.get("year"), source.get("month"), source.get("day"))
    candidate_date = (candidate.get("year"), candidate.get("month"), candidate.get("day"))
    if source_date[0] and candidate_date[0] and source_date[0] == candidate_date[0]:
        score += 0.12
        matched.append("year")
        if source_date[1] and source_date[1] == candidate_date[1]:
            score += 0.10
            matched.append("month")
        elif source_date[1] and candidate_date[1]:
            conflicts.append("month")
        if source_date[2] and source_date[2] == candidate_date[2]:
            score += 0.18
            matched.append("day")
        elif source_date[2] and candidate_date[2]:
            conflicts.append("day")

    source_country = normalized_country(source.get("country_normalized") or source.get("country_raw"))
    candidate_country = normalized_country(candidate.get("country"))
    if source_country and candidate_country:
        if source_country == candidate_country:
            score += 0.12
            matched.append("country")
        else:
            conflicts.append("country")

    location_score = max(
        similarity(source.get("location_raw"), candidate.get("location_public")),
        similarity(source.get("location_raw"), candidate.get("area")),
    )
    area_score = max(
        similarity(source.get("area_raw"), candidate.get("area")),
        similarity(source.get("area_raw"), candidate.get("location_public")),
    )
    activity_score = similarity(source.get("activity_raw"), candidate.get("activity"))
    species_score = similarity(source.get("shark_label_raw"), candidate.get("species_common"))
    score += 0.20 * location_score + 0.08 * area_score + 0.08 * activity_score + 0.07 * species_score
    if location_score >= 0.75:
        matched.append("location")
    if area_score >= 0.75:
        matched.append("area")
    if activity_score >= 0.75:
        matched.append("activity")
    if species_score >= 0.75:
        matched.append("species_label")

    fatality = source.get("fatality_claim")
    if fatality is not None and candidate.get("fatal") is not None:
        if int(fatality) == int(candidate["fatal"]):
            score += 0.05
            matched.append("fatality")
        else:
            conflicts.append("fatality")

    exact_date = all(source_date) and source_date == candidate_date
    if exact_date and source_country == candidate_country and location_score >= 0.90:
        status = "exact_candidate"
    elif score >= 0.58 and source_date[0] == candidate_date[0] and (location_score >= 0.55 or area_score >= 0.75):
        status = "likely_candidate"
    else:
        status = "no_candidate"
    return {
        "status": status,
        "score": round(score, 4),
        "matched_fields": matched,
        "conflicting_fields": conflicts,
        "candidate_record_id": candidate.get("record_id"),
        "candidate_canonical_key": candidate.get("canonical_key"),
        "candidate_source_name": candidate.get("source_name"),
        "candidate_source_record_id": candidate.get("source_record_id"),
        "candidate_date": candidate.get("date_text"),
        "candidate_location": candidate.get("location_public"),
    }


def attach_duplicate_review(records: list[dict[str, Any]], comparison_records: list[dict[str, Any]]) -> None:
    by_year: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for candidate in comparison_records:
        if candidate.get("year") is not None:
            by_year[int(candidate["year"])].append(candidate)

    within_groups: dict[str, list[str]] = defaultdict(list)
    for record in records:
        within_groups[record["within_source_match_key"]].append(record["source_record_id"])

    for record in records:
        year_candidates = by_year.get(record.get("year"), [])
        if record.get("month"):
            month_candidates = [
                candidate
                for candidate in year_candidates
                if candidate.get("month") in (None, record["month"])
            ]
        else:
            month_candidates = year_candidates
        candidates = [score_candidate(record, candidate) for candidate in month_candidates]
        candidates = [candidate for candidate in candidates if candidate["status"] != "no_candidate"]
        candidates.sort(key=lambda item: item["score"], reverse=True)
        repeated = within_groups[record["within_source_match_key"]]
        status = candidates[0]["status"] if candidates else "unmatched"
        record["duplicate_review"] = {
            "status": status,
            "auto_merged": False,
            "within_source_possible_duplicate_ids": [item for item in repeated if item != record["source_record_id"]],
            "cross_source_candidates": candidates[:5],
            "review_note": "Candidate matches require human review; no source row was merged or promoted.",
        }


def source_link_from_sharks_happen_record(record: dict[str, Any]) -> SourceLink:
    claims = ["source_incident_claim"]
    for field, claim in (
        ("incident_date_raw", "incident_date"),
        ("location_raw", "location"),
        ("activity_raw", "activity"),
        ("shark_label_raw", "source_species_label"),
        ("shark_size_raw", "source_shark_size"),
        ("injury_raw", "injury"),
    ):
        if record.get(field):
            claims.append(claim)
    return SourceLink(
        source_id=record["source_record_id"],
        source_type="sharks_happen_spreadsheet",
        source_name=SOURCE_NAME,
        source_ref=f"{record.get('source_workbook')} row {record.get('source_row_number')}",
        source_date=record.get("incident_date_raw"),
        source_title=SOURCE_NAME,
        source_rights_note=record.get("source_rights_note"),
        source_confidence="unreviewed",
        linked_claims=claims,
        public_citation_allowed=False,
        private_notes=f"Creator: {SOURCE_CREATOR}; channel: {SOURCE_CHANNEL}; duplicate review: {record.get('duplicate_review', {}).get('status', 'not_run')}",
    )


def persist_to_mongo(db: Any, records: list[dict[str, Any]], report: dict[str, Any], imported_at: datetime) -> dict[str, Any]:
    from app.mongodb import COLLECTIONS

    collection = db[COLLECTIONS["sharks_happen_sources"]]
    for record in records:
        document = dict(record)
        document["last_imported_at"] = imported_at.isoformat()
        document["ingest_source"] = "sharks_happen_workbook_import"
        collection.replace_one({"source_record_id": record["source_record_id"]}, document, upsert=True)
    db[COLLECTIONS["sharks_happen_import_reports"]].insert_one({**report, "visibility": "internal"})
    return {
        "enabled": True,
        "collection": COLLECTIONS["sharks_happen_sources"],
        "report_collection": COLLECTIONS["sharks_happen_import_reports"],
        "records_upserted": len(records),
        "reports_inserted": 1,
    }


def write_json(path: str | Path, payload: Any) -> None:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=True, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_duplicate_review_csv(path: str | Path, records: list[dict[str, Any]]) -> None:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "source_record_id",
        "source_count",
        "incident_date_raw",
        "victim_name_private",
        "location_raw",
        "area_raw",
        "country_raw",
        "activity_raw",
        "duplicate_status",
        "within_source_possible_duplicate_ids",
        "top_candidate_score",
        "top_candidate_record_id",
        "top_candidate_canonical_key",
        "top_candidate_source_name",
        "top_candidate_source_record_id",
        "top_candidate_date",
        "top_candidate_location",
        "matched_fields",
        "conflicting_fields",
    ]
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            review = record["duplicate_review"]
            candidates = review["cross_source_candidates"]
            top = candidates[0] if candidates else {}
            writer.writerow(
                {
                    "source_record_id": record["source_record_id"],
                    "source_count": record["source_count"],
                    "incident_date_raw": record.get("incident_date_raw"),
                    "victim_name_private": record.get("victim_name_private"),
                    "location_raw": record.get("location_raw"),
                    "area_raw": record.get("area_raw"),
                    "country_raw": record.get("country_raw"),
                    "activity_raw": record.get("activity_raw"),
                    "duplicate_status": review["status"],
                    "within_source_possible_duplicate_ids": ";".join(review["within_source_possible_duplicate_ids"]),
                    "top_candidate_score": top.get("score"),
                    "top_candidate_record_id": top.get("candidate_record_id"),
                    "top_candidate_canonical_key": top.get("candidate_canonical_key"),
                    "top_candidate_source_name": top.get("candidate_source_name"),
                    "top_candidate_source_record_id": top.get("candidate_source_record_id"),
                    "top_candidate_date": top.get("candidate_date"),
                    "top_candidate_location": top.get("candidate_location"),
                    "matched_fields": ";".join(top.get("matched_fields", [])),
                    "conflicting_fields": ";".join(top.get("conflicting_fields", [])),
                }
            )


def import_sharks_happen(
    input_path: str | Path,
    *,
    comparison_database: str | Path = DEFAULT_COMPARISON_DATABASE,
    staging_path: str | Path = DEFAULT_STAGING_PATH,
    report_path: str | Path = DEFAULT_REPORT_PATH,
    review_path: str | Path = DEFAULT_REVIEW_PATH,
    imported_at: datetime | None = None,
    mongo_database: Any | None = None,
) -> dict[str, Any]:
    timestamp = imported_at or utc_now()
    records = read_workbook(input_path, imported_at=timestamp)
    comparison_records = load_comparison_records(comparison_database)
    attach_duplicate_review(records, comparison_records)
    statuses = Counter(record["duplicate_review"]["status"] for record in records)
    within_source_rows = sum(bool(record["duplicate_review"]["within_source_possible_duplicate_ids"]) for record in records)
    warnings = Counter(warning for record in records for warning in record["normalization_warnings"])
    report = {
        "source_name": SOURCE_NAME,
        "source_creator": SOURCE_CREATOR,
        "source_channel": SOURCE_CHANNEL,
        "schema_version": SCHEMA_VERSION,
        "input_file": Path(input_path).name,
        "imported_at": timestamp.isoformat(),
        "records_written": len(records),
        "comparison_database": str(comparison_database),
        "duplicate_review_file": str(review_path),
        "comparison_records_considered": len(comparison_records),
        "duplicate_review_counts": dict(sorted(statuses.items())),
        "within_source_possible_duplicate_rows": within_source_rows,
        "normalization_warning_counts": dict(sorted(warnings.items())),
        "automatic_merges": 0,
        "automatic_registry_promotions": 0,
        "source_claims_are_confirmed_facts": False,
        "side_effects": dict(SIDE_EFFECTS),
    }
    if mongo_database is not None:
        report["mongo_persistence"] = persist_to_mongo(mongo_database, records, report, timestamp)
    staging = {
        **report,
        "records": records,
        "registry_source_links": [source_link_from_sharks_happen_record(record).model_dump(mode="json") for record in records],
    }
    write_json(staging_path, staging)
    write_json(report_path, report)
    write_duplicate_review_csv(review_path, records)
    return report


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import the local Sharks Happen Stats workbook as an attributed AI1SAD source.")
    parser.add_argument("--input", required=True, help="Path to Sharks Happen Stats.xlsx")
    parser.add_argument("--comparison-database", default=str(DEFAULT_COMPARISON_DATABASE))
    parser.add_argument("--staging", default=str(DEFAULT_STAGING_PATH))
    parser.add_argument("--report", default=str(DEFAULT_REPORT_PATH))
    parser.add_argument("--review", default=str(DEFAULT_REVIEW_PATH), help="CSV duplicate-review queue output path.")
    parser.add_argument("--mongo", action="store_true", help="Persist internal source records and the report to configured MongoDB.")
    parser.add_argument("--skip-mongo-indexes", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    mongo_database = None
    if args.mongo:
        from app.mongodb import ensure_sharks_happen_indexes, get_database

        mongo_database = get_database()
        if not args.skip_mongo_indexes:
            ensure_sharks_happen_indexes(mongo_database)
    report = import_sharks_happen(
        args.input,
        comparison_database=args.comparison_database,
        staging_path=args.staging,
        report_path=args.report,
        review_path=args.review,
        mongo_database=mongo_database,
    )
    print(json.dumps({"status": "ok", **report}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
